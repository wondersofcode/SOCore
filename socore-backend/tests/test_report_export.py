"""
Regression test for the Excel export: the Executive Summary's TOTAL ALERTS
KPI and the Alerts sheet row count must both match windowed_alerts() for the
same window — and must actually include alerts whose timestamp is one of the
non-naive shapes (Wazuh's real "+0000" offset, 'Z', etc.), not just 0.
"""
import io
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from app import report_export
from app.models import Alert, AlertStatus, Severity
from app.shift_summary import windowed_alerts

NOW = datetime.now(timezone.utc)


def _alert(id_: str, timestamp: str) -> Alert:
    return Alert(
        id=id_,
        timestamp=timestamp,
        severity=Severity.high,
        sourceIP="10.0.0.1",
        attackType="Brute Force",
        mitreId="T1110",
        mitreName="Brute Force",
        status=AlertStatus.new,
    )


def test_workbook_alert_count_matches_windowed_alerts_with_mixed_timestamp_shapes():
    alerts = [
        _alert("A1", (NOW - timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S")),
        _alert("A2", (NOW - timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M:%S.000+0000")),
        _alert("A3", (NOW - timedelta(hours=3)).strftime("%Y-%m-%dT%H:%M:%SZ")),
        _alert("A4", (NOW - timedelta(hours=30)).strftime("%Y-%m-%d %H:%M:%S")),  # outside 24h
    ]
    expected = len(windowed_alerts(alerts, 24))
    assert expected == 3  # sanity: the parser must actually accept all three in-window shapes

    workbook_bytes = report_export.build_workbook(
        alerts=alerts,
        cases=[],
        decisions=[],
        events=[],
        hours=24,
        summary_text="test",
        generated_at="2026-09-12 12:00:00",
        generated_by="tester",
    )
    wb = load_workbook(io.BytesIO(workbook_bytes))

    alerts_sheet = wb["Alerts"]
    # Sheet has a header row, one row per windowed alert, then a blank
    # spacer and a footer line — so count only rows whose ID column matches
    # a real alert id, not the raw row count.
    alert_ids = {a.id for a in alerts}
    data_rows = sum(
        1 for row in alerts_sheet.iter_rows(min_row=2, max_col=1, values_only=True)
        if row[0] in alert_ids
    )
    assert data_rows == expected

    exec_sheet = wb["Executive Summary"]
    total_alerts_value = None
    for row in exec_sheet.iter_rows():
        for cell in row:
            if cell.value == "TOTAL ALERTS":
                total_alerts_value = exec_sheet.cell(row=cell.row + 1, column=cell.column).value
    assert total_alerts_value == expected


def test_workbook_zero_real_alerts_does_not_fabricate_data():
    alerts = [_alert("A1", (NOW - timedelta(hours=30)).strftime("%Y-%m-%d %H:%M:%S"))]  # outside window
    workbook_bytes = report_export.build_workbook(
        alerts=alerts,
        cases=[],
        decisions=[],
        events=[],
        hours=24,
        summary_text="No alerts.",
        generated_at="2026-09-12 12:00:00",
        generated_by="tester",
    )
    wb = load_workbook(io.BytesIO(workbook_bytes))
    alerts_sheet = wb["Alerts"]
    assert alerts_sheet.cell(row=2, column=1).value == "No alerts were recorded during this shift."
    # No alert id anywhere on the sheet — nothing fabricated to fill the gap.
    assert not any(
        row[0] == "A1" for row in alerts_sheet.iter_rows(min_row=2, max_col=1, values_only=True)
    )
