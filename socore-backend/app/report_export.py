"""
Excel export of the shift report — the same window/data the AI Shift Summary
card shows, as a downloadable .xlsx with one sheet per data slice.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from . import shift_summary
from .models import Alert, Case, DecisionRecord

_HEADER_FONT = Font(bold=True)


def _write_table(ws: Worksheet, headers: list[str], rows: list[list]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _HEADER_FONT
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    # Auto-size columns to the widest cell (header or value), capped so one
    # long raw-log string can't blow a column out to an unusable width.
    for col_idx, header in enumerate(headers, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        widest = len(str(header))
        for row in rows:
            value = row[col_idx - 1]
            widest = max(widest, len(str(value)) if value is not None else 0)
        ws.column_dimensions[letter].width = min(widest + 2, 60)


def build_workbook(
    alerts: list[Alert],
    cases: list[Case],
    decisions: list[DecisionRecord],
    hours: int,
    summary_text: str,
    generated_at: str,
) -> bytes:
    windowed = shift_summary.windowed_alerts(alerts, hours)
    opened_cases = shift_summary.windowed_cases(cases, hours)
    windowed_ids = {a.id for a in windowed}
    windowed_decisions = [d for d in decisions if d.alertId in windowed_ids]

    critical_count = sum(1 for a in windowed if a.severity.value == "Critical")
    avg_risk = round(sum(a.riskScore for a in windowed) / len(windowed), 1) if windowed else 0

    wb = Workbook()

    # ── Summary ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "SOCore Shift Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Window: last {hours} hours"
    ws["A3"] = f"Generated: {generated_at}"
    ws["A5"] = "AI Shift Summary"
    ws["A5"].font = _HEADER_FONT
    ws["A6"] = summary_text
    ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A6:F10")
    ws.column_dimensions["A"].width = 100

    stats_row = 12
    ws.cell(row=stats_row, column=1, value="Metric").font = _HEADER_FONT
    ws.cell(row=stats_row, column=2, value="Value").font = _HEADER_FONT
    stats = [
        ("Total alerts in window", len(windowed)),
        ("Critical alerts", critical_count),
        ("Average risk score", avg_risk),
        ("Cases opened in window", len(opened_cases)),
        ("Decisions logged in window", len(windowed_decisions)),
    ]
    for i, (label, value) in enumerate(stats, start=stats_row + 1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["B"].width = 14

    # ── Alerts ───────────────────────────────────────────────────────────
    ws_alerts = wb.create_sheet("Alerts")
    _write_table(
        ws_alerts,
        ["ID", "Timestamp", "Severity", "Source IP", "Attack Type", "Risk Score", "Status", "Analyst"],
        [
            [a.id, a.timestamp, a.severity.value, a.sourceIP, a.attackType, a.riskScore, a.status.value, a.analyst]
            for a in windowed
        ],
    )

    # ── Cases ────────────────────────────────────────────────────────────
    ws_cases = wb.create_sheet("Cases")
    _write_table(
        ws_cases,
        ["ID", "Title", "Status", "Assigned To", "Created"],
        [[c.id, c.title, c.status.value, c.assignedTo, c.createdAt] for c in opened_cases],
    )

    # ── Decisions ────────────────────────────────────────────────────────
    ws_decisions = wb.create_sheet("Decisions")
    _write_table(
        ws_decisions,
        ["Alert ID", "Decision", "By", "At", "Reason"],
        [[d.alertId, d.status, d.by, d.at, d.reason] for d in windowed_decisions],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
